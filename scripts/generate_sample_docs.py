"""
Generate 10 synthetic sample documents for integration tests, worker tests, and E2E.

Writes to test/sample_docs/ (idempotent). This directory is the canonical sample set
for real-DB integration tests and E2E; no customer data is required.

Run: python scripts/generate_sample_docs.py
"""

from __future__ import annotations

import csv
from pathlib import Path

# Project root (parent of scripts/)
PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = PROJECT_ROOT / "test" / "sample_docs"


# ----- Content (deterministic, RAG-friendly) -----

TXT_MANUAL_PUMP = """Pump Priming Procedure

Equipment: Centrifugal pump model CP-200.

Step 1: Ensure the suction line is filled with liquid. Open the vent valve at the top of the pump casing until liquid flows without air bubbles.

Step 2: Close the vent valve. Verify the discharge valve is closed or throttled to prevent run-dry.

Step 3: Start the motor. Allow the pump to run for 30 seconds with the discharge closed to build pressure.

Step 4: Slowly open the discharge valve. Monitor the pressure gauge; normal operating range is 40-60 PSI for this model.

Step 5: Check for leaks at the mechanical seal and connections. If abnormal noise or vibration occurs, shut down and contact maintenance.
"""

TXT_SAFETY_CHECKLIST = """Safety Checklist - Field Service

Before starting any equipment work:

1. PPE: Wear safety glasses, gloves, and steel-toe boots. Hard hat required in designated areas.

2. Lockout/Tagout: Isolate energy sources and apply locks and tags. Verify zero energy before proceeding.

3. Ventilation: In confined spaces, ensure adequate ventilation and use a gas monitor if required.

4. First aid: Know the location of the nearest first aid kit and eyewash station.

5. Emergency contacts: Site supervisor and emergency number posted and accessible.

Do not proceed if any item cannot be verified. Escalate to supervisor.
"""

TXT_CONTACTS = """Munitor AI Service Contacts

Main office: 555-0100
After-hours dispatch: 555-0199

Technical support (in-house):
- Pumps and hydraulics: ext. 201
- Electrical and controls: ext. 202
- General equipment: ext. 203

Escalation: Site lead or project manager as per work order.
"""

MD_SOP_PUMP = """# Standard Operating Procedure: Pump Maintenance

## Scope
Applies to centrifugal pumps in the CP series installed at customer sites.

## Procedure

### Pre-maintenance
- Review work order and equipment history.
- Gather tools and replacement parts (seal kit, gaskets).
- Perform lockout/tagout per site requirements.

### Inspection
- Check mechanical seal for leaks.
- Inspect coupling alignment and wear.
- Record motor amperage and pump pressure.

### Post-maintenance
- Restore power and verify operation.
- Complete work order and update CMMS.
"""

MD_SOP_SAFETY = """# Safety SOP: Confined Space Entry

## Requirements
- Permit required for confined space entry.
- Attendant must be posted outside.
- Gas monitor must show safe O2 and no LEL.

## Steps
1. Obtain permit from site safety.
2. Test atmosphere at entry and at depth.
3. Use harness and retrieval line.
4. Attendant maintains constant communication.
"""

MD_README = """# Equipment Manual Summary

## Pump CP-200
Centrifugal pump, max head 60 m, max flow 200 L/min. Priming procedure documented in separate SOP.

## Maintenance Interval
- Seal inspection: every 6 months.
- Bearing lubrication: every 3 months.
"""

CSV_PARTS = """part_id,description,quantity_unit
P-001,Mechanical seal kit CP-200,each
P-002,Impeller CP-200,each
P-003,Gasket set discharge,each
P-004,Motor 5HP 3-phase,each
P-005,Bearing set,each
"""

CSV_STEPS = """step,action,notes
1,Isolate pump,Lockout electrical and close valves
2,Remove coupling guard,Set aside hardware
3,Drain pump,Use catch pan
4,Replace seal,Follow seal kit instructions
5,Reassemble and test,Verify no leaks
"""


def write_txt(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


def write_md(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


def write_csv(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


def write_docx(path: Path, content: str) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    for line in content.strip().split("\n"):
        line = line.strip()
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(6)
    doc.save(str(path))


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # 3 TXT
    write_txt(OUTPUT_DIR / "manual_pump_01.txt", TXT_MANUAL_PUMP)
    write_txt(OUTPUT_DIR / "safety_checklist_02.txt", TXT_SAFETY_CHECKLIST)
    write_txt(OUTPUT_DIR / "contacts_03.txt", TXT_CONTACTS)

    # 3 MD
    write_md(OUTPUT_DIR / "sop_pump_01.md", MD_SOP_PUMP)
    write_md(OUTPUT_DIR / "sop_safety_02.md", MD_SOP_SAFETY)
    write_md(OUTPUT_DIR / "readme_manual_03.md", MD_README)

    # 2 CSV
    write_csv(OUTPUT_DIR / "parts_01.csv", CSV_PARTS)
    write_csv(OUTPUT_DIR / "steps_02.csv", CSV_STEPS)

    # 2 DOCX (use same content as TXT/MD for consistency)
    write_docx(OUTPUT_DIR / "procedure_pump_01.docx", TXT_MANUAL_PUMP)
    write_docx(OUTPUT_DIR / "procedure_safety_02.docx", "# Safety Checklist\n\n" + TXT_SAFETY_CHECKLIST)

    print(f"Generated 10 sample documents in {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
